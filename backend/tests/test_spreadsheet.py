"""Spreadsheet (CSV/XLSX) ingestion + the native SpreadsheetEngine + grounding.

All offline: files are generated in-memory and the mock structuring provider grounds
against the parsed sheet, so no OCR/LLM deps are needed.
"""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.models import Document
from app.pipeline.ocr import available_engines, get_engine
from app.pipeline.ocr.spreadsheet import SpreadsheetEngine
from sqlmodel import Session

from app.db import engine as db_engine


def _xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # drop the default sheet
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client: TestClient, name: str, data: bytes, doc_type: str | None = None) -> dict:
    files = {"file": (name, data, "application/octet-stream")}
    form = {"doc_type": doc_type} if doc_type else None
    resp = client.post("/documents", files=files, data=form)
    assert resp.status_code == 201, resp.text
    return resp.json()


# --- ingestion ---------------------------------------------------------------


def test_upload_csv_normalizes_to_one_page():
    with TestClient(app) as client:
        detail = _upload(client, "data.csv", b"Vendor,Amount\nAcme Corp,1234.56\n")
        assert detail["mime"] == "text/csv"
        assert detail["page_count"] == 1

        sheets = client.get(f"/files/{detail['id']}/sheets.json").json()
        assert sheets[0]["name"] == "Sheet1"
        assert sheets[0]["rows"] == [["Vendor", "Amount"], ["Acme Corp", "1234.56"]]


def test_upload_xlsx_one_page_per_sheet():
    with TestClient(app) as client:
        data = _xlsx_bytes({"Invoices": [["Vendor", "Total"]], "Notes": [["n", "hi"]]})
        detail = _upload(client, "book.xlsx", data)
        assert detail["mime"].endswith("spreadsheetml.sheet")
        assert detail["page_count"] == 2  # one page per sheet

        sheets = client.get(f"/files/{detail['id']}/sheets.json").json()
        assert [s["name"] for s in sheets] == ["Invoices", "Notes"]


def test_unsupported_extension_still_rejected():
    with TestClient(app) as client:
        resp = client.post("/documents", files={"file": ("x.zip", b"PK\x03\x04", "application/zip")})
        assert resp.status_code == 415


# --- engine ------------------------------------------------------------------


def test_spreadsheet_engine_registered():
    with Session(db_engine) as session:
        assert "spreadsheet" in available_engines(session)
        assert isinstance(get_engine("spreadsheet", session), SpreadsheetEngine)


def test_engine_emits_one_cell_block_with_grid_coords():
    with TestClient(app) as client:
        data = _xlsx_bytes({"S": [["Vendor", "Total"], ["Acme Corp", 1234.56], ["Beta", 5.0]]})
        detail = _upload(client, "book.xlsx", data)

        # Any requested engine is overridden to the spreadsheet engine.
        resp = client.post(f"/documents/{detail['id']}/ocr", params={"engine": "docling"})
        assert resp.status_code == 200, resp.text
        result = resp.json()
        assert result["engine_name"] == "spreadsheet"
        assert result["table_count"] == 1

        blocks = result["pages"][0]["blocks"]
        # One block per non-empty cell (2x2 header/data + 2 = 6 cells).
        assert len(blocks) == 6
        by_text = {b["text"]: b for b in blocks}
        # Whole-number floats render cleanly; bbox = (col, row, col+1, row+1).
        assert by_text["5"]["bbox"] == [1.0, 2.0, 2.0, 3.0]
        assert by_text["Vendor"]["bbox"] == [0.0, 0.0, 1.0, 1.0]
        assert all(b["label"] == "cell" for b in blocks)

        # Table markdown carries the cell values for downstream grounding.
        assert "1234.56" in result["pages"][0]["tables"][0]["markdown"]


def test_prescan_skipped_for_spreadsheet():
    with TestClient(app) as client:
        detail = _upload(client, "data.csv", b"a,b\n1,2\n")
        resp = client.post(f"/documents/{detail['id']}/prescan")
        assert resp.status_code == 200, resp.text
        report = resp.json()
        assert report["verdict"] == "pass"
        assert report["pages"] == []
        assert report["preprocess_applied"] is False


def test_truncation_surfaces_as_warning():
    with TestClient(app) as client:
        rows = [[f"r{i}"] for i in range(600)]  # exceeds MAX_SHEET_ROWS (500)
        data = _xlsx_bytes({"Big": rows})
        detail = _upload(client, "big.xlsx", data)
        result = client.post(f"/documents/{detail['id']}/ocr").json()
        assert any("truncated" in w for w in result["warnings"])


# --- grounding (mock structuring provider grounds into the parsed sheet) ------


def test_structuring_grounds_fields_to_cells():
    """The mock invoice provider's spans are placed in the sheet so grounding resolves."""
    with TestClient(app) as client:
        # Cells contain exactly what the mock invoice provider extracts, so str.find
        # anchors each span in the sheet's table markdown.
        data = _xlsx_bytes(
            {"Invoice": [["Vendor", "MOCK INVOICE"], ["Ref", "page 1"], ["Total", "$1,234.56"]]}
        )
        detail = _upload(client, "invoice.xlsx", data, doc_type="invoice")
        doc_id = detail["id"]

        assert client.post(f"/documents/{doc_id}/prescan").status_code == 200
        assert client.post(f"/documents/{doc_id}/ocr").status_code == 200
        struct = client.post(
            f"/documents/{doc_id}/structure", params={"provider": "mock", "doc_type": "invoice"}
        )
        assert struct.status_code == 200, struct.text

        gm = struct.json()["grounding_map"]
        # Vendor/invoice_no/total live in cells -> grounded to the (single) sheet page.
        assert gm["vendor"]["page"] == 1
        assert gm["vendor"]["alignment"] == "exact"
        assert gm["total"]["snippet"] == "$1,234.56"
        assert gm["total"]["page"] == 1


def test_engine_run_offline_with_missing_sheets_json():
    """A doc with no sheets.json degrades to an empty result rather than raising."""
    doc = Document(id="no-sheets", filename="x.xlsx", mime="text/csv", page_count=0)
    result = SpreadsheetEngine().run(doc)
    assert result.engine_name == "spreadsheet"
    assert result.pages == []
